"""Safely import supplier/product assignments and managed order templates.

The source workbook is expected to contain the sheets produced by the supplier
reconciliation process:

* ``Found Items``: one confirmed supplier per SKU.
* ``Conflicts``: semicolon-separated supplier candidates.  These are treated as
  valid multi-vendor relationships, not as errors.

The command is a dry run unless ``--apply`` is supplied.  Apply mode creates a
verified ``pg_dump`` checkpoint before opening the write transaction.  Database
credentials are read from environment variables or an interactive password
prompt; they are never stored in this file.
"""

from __future__ import annotations

import argparse
import getpass
import os
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import psycopg2
from openpyxl import load_workbook
from psycopg2.extras import execute_values


MANAGED_TEMPLATE_MARKER = "[supplier-product-import:v1]"
REQUIRED_TABLES = {
    "containermgmt": {
        "products",
        "product_categories",
        "supplier",
        "product_suppliers",
        "order_templates",
        "order_template_items",
    },
    "usercredentials": {"organisations"},
}


@dataclass(frozen=True)
class ProductSource:
    sku: str
    name: str
    category: str | None
    brand: str | None
    suppliers: tuple[str, ...]
    source_kind: str


def clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def lookup_key(value: str) -> str:
    return clean_text(value).casefold()  # type: ignore[union-attr]


def iter_data_rows(sheet, header_row: int = 4) -> Iterable[dict[str, object]]:
    headers = [clean_text(cell.value) for cell in sheet[header_row]]
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in values):
            continue
        yield {
            header: values[index]
            for index, header in enumerate(headers)
            if header is not None and index < len(values)
        }


def load_source(path: Path) -> tuple[dict[str, ProductSource], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    required_sheets = {"Found Items", "Conflicts"}
    missing_sheets = required_sheets.difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Workbook is missing sheets: {sorted(missing_sheets)}")

    products: dict[str, ProductSource] = {}
    stats = Counter()

    def add_product(
        row: dict[str, object], suppliers: list[str], source_kind: str
    ) -> None:
        sku = clean_text(row.get("Item Code / SKU"))
        name = clean_text(row.get("Product Name"))
        if not sku or not name:
            raise ValueError(f"{source_kind} row is missing SKU or product name: {row}")
        cleaned_suppliers = tuple(dict.fromkeys(clean_text(v) for v in suppliers if clean_text(v)))
        if not cleaned_suppliers:
            raise ValueError(f"{source_kind} row {sku!r} has no supplier")

        key = lookup_key(sku)
        candidate = ProductSource(
            sku=sku,
            name=name,
            category=clean_text(row.get("Category")),
            brand=clean_text(row.get("Manufacturer / Brand")),
            suppliers=cleaned_suppliers,
            source_kind=source_kind,
        )
        existing = products.get(key)
        if existing and existing != candidate:
            raise ValueError(f"Duplicate SKU {sku!r} has inconsistent source data")
        products[key] = candidate
        stats[source_kind] += 1

    for row in iter_data_rows(workbook["Found Items"]):
        supplier = clean_text(row.get("Supplier"))
        add_product(row, [supplier] if supplier else [], "confirmed")

    for row in iter_data_rows(workbook["Conflicts"]):
        candidates = clean_text(row.get("Supplier Candidates"))
        suppliers = [part for part in (candidates or "").split(";") if clean_text(part)]
        add_product(row, suppliers, "multi_vendor")

    stats["products"] = len(products)
    stats["supplier_links"] = sum(len(item.suppliers) for item in products.values())
    stats["suppliers"] = len(
        {lookup_key(name) for item in products.values() for name in item.suppliers}
    )
    limits = {
        "sku": (100, (item.sku for item in products.values())),
        "product name": (255, (item.name for item in products.values())),
        "category": (100, (item.category for item in products.values() if item.category)),
        "brand": (100, (item.brand for item in products.values() if item.brand)),
        "supplier": (
            255,
            (name for item in products.values() for name in item.suppliers),
        ),
    }
    for label, (maximum, values) in limits.items():
        too_long = [value for value in values if len(value) > maximum]
        if too_long:
            raise ValueError(
                f"Workbook contains {label} values longer than {maximum}: {too_long[:5]}"
            )
    return products, dict(stats)


def connect(args: argparse.Namespace, supplied_password: str | None = None):
    password = supplied_password or os.getenv("DB_PASSWORD") or os.getenv("PGPASSWORD")
    if not password:
        password = getpass.getpass(f"Password for {args.user}@{args.host}:{args.port}: ")
    return psycopg2.connect(
        host=args.host,
        port=args.port,
        dbname=args.database,
        user=args.user,
        password=password,
        sslmode=args.sslmode,
        connect_timeout=args.connect_timeout,
        application_name="freightlens_supplier_product_import",
        options=f"-c statement_timeout={args.statement_timeout_ms}",
    ), password


def verify_schema(cursor) -> None:
    for schema_name, table_names in REQUIRED_TABLES.items():
        cursor.execute(
            """
            SELECT table_name
            FROM information_schema.tables
            WHERE table_schema = %s AND table_name = ANY(%s)
            """,
            (schema_name, list(table_names)),
        )
        found = {row[0] for row in cursor.fetchall()}
        missing = table_names.difference(found)
        if missing:
            raise RuntimeError(f"Database is missing {schema_name} tables: {sorted(missing)}")


def resolve_org_id(cursor, requested_org_id: int | None, sku_keys: list[str]) -> int:
    cursor.execute(
        """
        SELECT id, name, display_name
        FROM usercredentials.organisations
        WHERE is_active = TRUE
        ORDER BY id
        """
    )
    active_orgs = cursor.fetchall()
    active_ids = {row[0] for row in active_orgs}
    if requested_org_id is not None:
        if requested_org_id not in active_ids:
            raise ValueError(f"Organisation {requested_org_id} is not active or does not exist")
        return requested_org_id

    cursor.execute(
        """
        SELECT org_id, COUNT(*)
        FROM containermgmt.products
        WHERE is_deleted = FALSE AND lower(btrim(sku)) = ANY(%s)
        GROUP BY org_id
        ORDER BY COUNT(*) DESC, org_id
        """,
        (sku_keys,),
    )
    matches = cursor.fetchall()
    if len(matches) == 1:
        return matches[0][0]
    if not matches and len(active_orgs) == 1:
        return active_orgs[0][0]

    org_description = ", ".join(
        f"{row[0]}:{row[2] or row[1]}" for row in active_orgs
    )
    raise ValueError(
        "Cannot infer one target organisation. Pass --org-id. "
        f"Active organisations: {org_description or 'none'}"
    )


def inspect_database(cursor, products: dict[str, ProductSource], org_id: int) -> dict[str, object]:
    sku_keys = list(products)
    cursor.execute(
        """
        SELECT lower(btrim(sku)), COUNT(*), array_agg(id ORDER BY id)
        FROM containermgmt.products
        WHERE org_id = %s AND is_deleted = FALSE AND lower(btrim(sku)) = ANY(%s)
        GROUP BY lower(btrim(sku))
        """,
        (org_id, sku_keys),
    )
    product_matches = {row[0]: (row[1], row[2]) for row in cursor.fetchall()}
    duplicate_skus = {key: ids for key, (count, ids) in product_matches.items() if count > 1}
    if duplicate_skus:
        sample = dict(list(duplicate_skus.items())[:10])
        raise ValueError(f"Duplicate active product SKUs in organisation {org_id}: {sample}")

    supplier_names = sorted(
        {name for product in products.values() for name in product.suppliers}, key=str.casefold
    )
    cursor.execute(
        """
        SELECT supplier_id, name, org_id, is_shared
        FROM containermgmt.supplier
        WHERE is_deleted = FALSE AND (is_shared = TRUE OR org_id = %s)
        ORDER BY is_shared DESC, supplier_id
        """,
        (org_id,),
    )
    suppliers_by_key: dict[str, list[tuple]] = defaultdict(list)
    for row in cursor.fetchall():
        if clean_text(row[1]):
            suppliers_by_key[lookup_key(row[1])].append(row)
    ambiguous = {
        name: suppliers_by_key[lookup_key(name)]
        for name in supplier_names
        if len(suppliers_by_key.get(lookup_key(name), [])) > 1
    }
    if ambiguous:
        raise ValueError(f"Ambiguous visible suppliers: {dict(list(ambiguous.items())[:10])}")

    found_suppliers = {
        lookup_key(name): suppliers_by_key[lookup_key(name)][0]
        for name in supplier_names
        if suppliers_by_key.get(lookup_key(name))
    }
    missing_suppliers = [name for name in supplier_names if lookup_key(name) not in found_suppliers]

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM containermgmt.order_templates
        WHERE org_id = %s AND is_deleted = FALSE AND notes LIKE %s
        """,
        (org_id, f"%{MANAGED_TEMPLATE_MARKER}%"),
    )
    managed_templates = cursor.fetchone()[0]

    return {
        "org_id": org_id,
        "existing_products": len(product_matches),
        "new_products": len(products) - len(product_matches),
        "existing_suppliers": len(found_suppliers),
        "missing_suppliers": missing_suppliers,
        "managed_templates": managed_templates,
    }


def find_pg_dump(explicit_path: str | None) -> str:
    candidates = [explicit_path, shutil.which("pg_dump")]
    candidates.extend(
        str(path)
        for path in Path("C:/Program Files/PostgreSQL").glob("*/bin/pg_dump.exe")
    )
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate))
    raise FileNotFoundError("pg_dump was not found; pass --pg-dump")


def create_checkpoint(args: argparse.Namespace, password: str, org_id: int) -> Path:
    checkpoint_dir = Path(args.checkpoint_dir).resolve()
    checkpoint_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    target = checkpoint_dir / f"{args.database}-pre-supplier-product-import-org{org_id}-{timestamp}.dump"
    pg_dump = find_pg_dump(args.pg_dump)
    env = os.environ.copy()
    env["PGPASSWORD"] = password
    command = [
        pg_dump,
        "--host", args.host,
        "--port", str(args.port),
        "--username", args.user,
        "--dbname", args.database,
        "--format", "custom",
        "--file", str(target),
        "--no-password",
    ]
    if args.sslmode:
        env["PGSSLMODE"] = args.sslmode
    try:
        completed = subprocess.run(
            command,
            env=env,
            check=False,
            capture_output=True,
            text=True,
            timeout=args.checkpoint_timeout,
        )
    except Exception:
        target.unlink(missing_ok=True)
        raise
    if completed.returncode != 0 or not target.exists() or target.stat().st_size == 0:
        target.unlink(missing_ok=True)
        detail = (completed.stderr or completed.stdout or "unknown pg_dump failure").strip()
        raise RuntimeError(f"Checkpoint failed: {detail}")
    return target


def upsert_data(cursor, products: dict[str, ProductSource], org_id: int) -> dict[str, int]:
    counts = Counter()
    supplier_names = sorted(
        {name for product in products.values() for name in product.suppliers}, key=str.casefold
    )
    category_names = sorted(
        {p.category for p in products.values() if p.category}, key=str.casefold
    )

    cursor.execute("SELECT pg_advisory_xact_lock(%s, %s)", (org_id, 918274))
    cursor.execute("CREATE TEMP TABLE import_products (sku_key TEXT PRIMARY KEY, sku TEXT, name TEXT, category_key TEXT, category_name TEXT, brand TEXT) ON COMMIT DROP")
    execute_values(cursor, "INSERT INTO import_products VALUES %s", [
        (key, item.sku, item.name, lookup_key(item.category) if item.category else None, item.category, item.brand)
        for key, item in products.items()
    ], page_size=1000)
    cursor.execute("CREATE TEMP TABLE import_suppliers (supplier_key TEXT PRIMARY KEY, supplier_name TEXT NOT NULL) ON COMMIT DROP")
    execute_values(cursor, "INSERT INTO import_suppliers VALUES %s", [(lookup_key(name), name) for name in supplier_names], page_size=1000)
    cursor.execute("CREATE TEMP TABLE import_product_suppliers (sku_key TEXT NOT NULL, supplier_key TEXT NOT NULL, PRIMARY KEY (sku_key, supplier_key)) ON COMMIT DROP")
    execute_values(cursor, "INSERT INTO import_product_suppliers VALUES %s", [
        (sku_key, lookup_key(name)) for sku_key, item in products.items() for name in item.suppliers
    ], page_size=1000)

    cursor.execute("""
        INSERT INTO containermgmt.supplier
            (org_id, is_shared, name, is_active, default_currency, is_deleted, created_at, updated_at)
        SELECT %s, FALSE, src.supplier_name, TRUE, 'USD', FALSE, NOW(), NOW()
        FROM import_suppliers src
        WHERE NOT EXISTS (
            SELECT 1 FROM containermgmt.supplier s
            WHERE s.is_deleted = FALSE AND (s.is_shared = TRUE OR s.org_id = %s)
              AND lower(btrim(s.name)) = src.supplier_key
        )
    """, (org_id, org_id))
    counts["suppliers_created"] = cursor.rowcount

    cursor.execute("""
        INSERT INTO containermgmt.product_categories
            (org_id, name, is_shared, is_subcategory, is_deleted, created_at, updated_at)
        SELECT %s, src.category_name, FALSE, FALSE, FALSE, NOW(), NOW()
        FROM (SELECT DISTINCT category_key, category_name FROM import_products WHERE category_key IS NOT NULL) src
        WHERE NOT EXISTS (
            SELECT 1 FROM containermgmt.product_categories c
            WHERE c.org_id = %s AND c.is_deleted = FALSE
              AND lower(btrim(c.name)) = src.category_key
        )
    """, (org_id, org_id))
    counts["categories_created"] = cursor.rowcount

    cursor.execute("""
        UPDATE containermgmt.products p
        SET name = src.name,
            category_id = cat.id,
            brand = src.brand,
            status = 'active',
            updated_at = NOW()
        FROM import_products src
        LEFT JOIN containermgmt.product_categories cat
          ON cat.org_id = %s AND cat.is_deleted = FALSE
         AND lower(btrim(cat.name)) = src.category_key
        WHERE p.org_id = %s AND p.is_deleted = FALSE
          AND lower(btrim(p.sku)) = src.sku_key
    """, (org_id, org_id))
    counts["products_updated"] = cursor.rowcount

    cursor.execute("""
        INSERT INTO containermgmt.products
            (org_id, is_shared, code, sku, name, category_id, brand, unit,
             status, currency, current_stock, min_stock_quantity,
             is_deleted, created_at, updated_at)
        SELECT %s, FALSE, src.sku, src.sku, src.name, cat.id, src.brand, 'PCS',
               'active', 'USD', 0, 0, FALSE, NOW(), NOW()
        FROM import_products src
        LEFT JOIN containermgmt.product_categories cat
          ON cat.org_id = %s AND cat.is_deleted = FALSE
         AND lower(btrim(cat.name)) = src.category_key
        WHERE NOT EXISTS (
            SELECT 1 FROM containermgmt.products p
            WHERE p.org_id = %s AND p.is_deleted = FALSE
              AND lower(btrim(p.sku)) = src.sku_key
        )
    """, (org_id, org_id, org_id))
    counts["products_created"] = cursor.rowcount

    cursor.execute("""
        CREATE TEMP TABLE import_supplier_map ON COMMIT DROP AS
        SELECT src.supplier_key, min(s.supplier_id) AS supplier_id, min(s.name) AS supplier_name
        FROM import_suppliers src
        JOIN containermgmt.supplier s
          ON s.is_deleted = FALSE AND (s.is_shared = TRUE OR s.org_id = %s)
         AND lower(btrim(s.name)) = src.supplier_key
        GROUP BY src.supplier_key
    """, (org_id,))
    cursor.execute("""
        CREATE TEMP TABLE import_product_map ON COMMIT DROP AS
        SELECT src.sku_key, min(p.id) AS product_id, min(p.sku) AS sku,
               min(p.name) AS name
        FROM import_products src
        JOIN containermgmt.products p
          ON p.org_id = %s AND p.is_deleted = FALSE
         AND lower(btrim(p.sku)) = src.sku_key
        GROUP BY src.sku_key
    """, (org_id,))
    cursor.execute("""
        CREATE TEMP TABLE import_resolved_links ON COMMIT DROP AS
        SELECT pm.sku_key, pm.product_id, pm.sku, pm.name,
               sm.supplier_key, sm.supplier_id, sm.supplier_name,
               (count(*) OVER (PARTITION BY ips.sku_key) = 1) AS is_default
        FROM import_product_suppliers ips
        JOIN import_product_map pm USING (sku_key)
        JOIN import_supplier_map sm USING (supplier_key)
    """)

    cursor.execute("""
        SELECT p.sku, s.supplier_name, count(*)
        FROM containermgmt.product_suppliers ps
        JOIN import_product_map p ON p.product_id = ps.product_id
        JOIN import_supplier_map s ON s.supplier_id = ps.supplier_id
        WHERE ps.org_id = %s
        GROUP BY p.sku, s.supplier_name HAVING count(*) > 1
        LIMIT 10
    """, (org_id,))
    duplicates = cursor.fetchall()
    if duplicates:
        raise ValueError(f"Duplicate product-supplier links prevent safe import: {duplicates}")

    cursor.execute("""
        UPDATE containermgmt.product_suppliers ps
        SET factory_code = src.sku, vendor_product_name = src.name,
            is_default = src.is_default, is_deleted = FALSE, deleted_at = NULL,
            updated_at = NOW()
        FROM import_resolved_links src
        WHERE ps.org_id = %s AND ps.product_id = src.product_id
          AND ps.supplier_id = src.supplier_id
    """, (org_id,))
    counts["supplier_links_updated"] = cursor.rowcount
    cursor.execute("""
        INSERT INTO containermgmt.product_suppliers
            (org_id, product_id, supplier_id, factory_code, vendor_product_name,
             currency, is_default, is_deleted, created_at, updated_at)
        SELECT %s, src.product_id, src.supplier_id, src.sku, src.name,
               'USD', src.is_default, FALSE, NOW(), NOW()
        FROM import_resolved_links src
        WHERE NOT EXISTS (
            SELECT 1 FROM containermgmt.product_suppliers ps
            WHERE ps.org_id = %s AND ps.product_id = src.product_id
              AND ps.supplier_id = src.supplier_id
        )
    """, (org_id, org_id))
    counts["supplier_links_created"] = cursor.rowcount

    cursor.execute("""
        UPDATE containermgmt.products p
        SET default_supplier_id = CASE
              WHEN totals.supplier_count = 1 THEN single.supplier_id
              WHEN p.default_supplier_id = ANY(totals.supplier_ids) THEN p.default_supplier_id
              ELSE NULL
            END,
            updated_at = NOW()
        FROM (
            SELECT product_id, count(*) AS supplier_count,
                   array_agg(supplier_id) AS supplier_ids,
                   min(supplier_id) AS supplier_id
            FROM import_resolved_links GROUP BY product_id
        ) totals
        LEFT JOIN LATERAL (
            SELECT supplier_id FROM import_resolved_links x
            WHERE x.product_id = totals.product_id AND x.is_default
            LIMIT 1
        ) single ON TRUE
        WHERE p.id = totals.product_id
    """)

    cursor.execute("""
        CREATE TEMP TABLE import_template_map ON COMMIT DROP AS
        SELECT sm.supplier_id, sm.supplier_name,
               min(ot.id) FILTER (WHERE ot.id IS NOT NULL) AS template_id,
               count(ot.id) AS template_count
        FROM import_supplier_map sm
        LEFT JOIN containermgmt.order_templates ot
          ON ot.org_id = %s AND ot.supplier_id = sm.supplier_id
         AND ot.is_deleted = FALSE AND ot.notes LIKE %s
        GROUP BY sm.supplier_id, sm.supplier_name
    """, (org_id, f"%{MANAGED_TEMPLATE_MARKER}%"))
    cursor.execute("SELECT supplier_name, template_count FROM import_template_map WHERE template_count > 1")
    duplicate_templates = cursor.fetchall()
    if duplicate_templates:
        raise ValueError(f"Multiple managed templates exist for suppliers: {duplicate_templates}")

    notes = f"{MANAGED_TEMPLATE_MARKER} Generated from reconciled supplier workbook."
    cursor.execute("""
        INSERT INTO containermgmt.order_templates
            (org_id, name, description, supplier_id, company, freight_type,
             visibility, notes, is_deleted, created_at, updated_at)
        SELECT %s, sm.supplier_name || ' - Product Order Template',
               'Products available from ' || sm.supplier_name, sm.supplier_id,
               sm.supplier_name, 'Sea Freight', 'org', %s, FALSE, NOW(), NOW()
        FROM import_template_map sm
        WHERE sm.template_id IS NULL
        RETURNING id, supplier_id
    """, (org_id, notes))
    new_templates = cursor.fetchall()
    counts["templates_created"] = len(new_templates)
    if new_templates:
        execute_values(
            cursor,
            "UPDATE import_template_map AS target SET template_id = source.template_id FROM (VALUES %s) AS source(supplier_id, template_id) WHERE target.supplier_id = source.supplier_id",
            [(supplier_id, template_id) for template_id, supplier_id in new_templates],
            page_size=1000,
        )
    cursor.execute("""
        UPDATE containermgmt.order_templates ot
        SET name = sm.supplier_name || ' - Product Order Template',
            description = 'Products available from ' || sm.supplier_name,
            company = sm.supplier_name, freight_type = 'Sea Freight',
            visibility = 'org', notes = %s, updated_at = NOW()
        FROM import_template_map sm
        WHERE ot.id = sm.template_id AND sm.template_id IS NOT NULL
          AND sm.template_count > 0
    """, (notes,))
    counts["templates_updated"] = cursor.rowcount

    cursor.execute("""
        UPDATE containermgmt.order_template_items oti
        SET is_deleted = TRUE, deleted_at = NOW(), updated_at = NOW()
        FROM import_template_map tm
        WHERE oti.template_id = tm.template_id AND oti.is_deleted = FALSE
    """)
    cursor.execute("""
        CREATE TEMP TABLE import_template_items ON COMMIT DROP AS
        SELECT tm.template_id, l.product_id, l.sku, l.name,
               row_number() OVER (PARTITION BY tm.template_id ORDER BY l.sku_key) AS sort_order
        FROM import_resolved_links l
        JOIN import_template_map tm USING (supplier_id)
    """)
    cursor.execute("SELECT template_id, product_id, sku, name, sort_order FROM import_template_items ORDER BY template_id, sort_order")
    template_items = cursor.fetchall()
    execute_values(cursor, """
        INSERT INTO containermgmt.order_template_items
            (template_id, product_id, item_code, description, default_quantity,
             unit, currency, sort_order, is_deleted, created_at, updated_at)
        VALUES %s
    """, [
        (
            template_id, product_id, sku, name, 1, "PCS", "USD",
            int(sort_order), False, datetime.now(timezone.utc), datetime.now(timezone.utc)
        )
        for template_id, product_id, sku, name, sort_order in template_items
    ], page_size=1000)
    counts["template_items_created"] = len(template_items)
    return dict(counts)


def verify_applied(cursor, products: dict[str, ProductSource], org_id: int) -> dict[str, int]:
    expected_links = sum(len(item.suppliers) for item in products.values())
    expected_suppliers = len({lookup_key(n) for p in products.values() for n in p.suppliers})
    sku_keys = list(products)
    cursor.execute(
        """
        CREATE TEMP TABLE expected_supplier_product_import (
            sku_key TEXT NOT NULL,
            supplier_key TEXT NOT NULL,
            PRIMARY KEY (sku_key, supplier_key)
        ) ON COMMIT DROP
        """
    )
    execute_values(
        cursor,
        """
        INSERT INTO expected_supplier_product_import (sku_key, supplier_key)
        VALUES %s
        """,
        [
            (sku_key, lookup_key(supplier_name))
            for sku_key, item in products.items()
            for supplier_name in item.suppliers
        ],
        page_size=1000,
    )
    cursor.execute(
        """
        SELECT COUNT(*)
        FROM containermgmt.products
        WHERE org_id = %s AND is_deleted = FALSE AND lower(btrim(sku)) = ANY(%s)
        """,
        (org_id, sku_keys),
    )
    product_count = cursor.fetchone()[0]
    cursor.execute(
        """
        SELECT COUNT(*)
        FROM expected_supplier_product_import expected
        JOIN containermgmt.products p
          ON p.org_id = %s AND p.is_deleted = FALSE
         AND lower(btrim(p.sku)) = expected.sku_key
        JOIN containermgmt.supplier s
          ON s.is_deleted = FALSE AND (s.is_shared = TRUE OR s.org_id = %s)
         AND lower(btrim(s.name)) = expected.supplier_key
        JOIN containermgmt.product_suppliers ps
          ON ps.org_id = %s AND ps.is_deleted = FALSE
         AND ps.product_id = p.id AND ps.supplier_id = s.supplier_id
        """,
        (org_id, org_id, org_id),
    )
    link_count = cursor.fetchone()[0]
    cursor.execute(
        """
        SELECT COUNT(DISTINCT ot.supplier_id), COUNT(DISTINCT oti.id)
        FROM containermgmt.order_templates ot
        JOIN containermgmt.supplier s
          ON s.supplier_id = ot.supplier_id AND s.is_deleted = FALSE
        JOIN expected_supplier_product_import expected_supplier
          ON expected_supplier.supplier_key = lower(btrim(s.name))
        LEFT JOIN containermgmt.order_template_items oti
          ON oti.template_id = ot.id AND oti.is_deleted = FALSE
        LEFT JOIN containermgmt.products p
          ON p.id = oti.product_id AND p.is_deleted = FALSE
        LEFT JOIN expected_supplier_product_import expected_item
          ON expected_item.supplier_key = lower(btrim(s.name))
         AND expected_item.sku_key = lower(btrim(p.sku))
        WHERE ot.org_id = %s AND ot.is_deleted = FALSE AND ot.notes LIKE %s
          AND (oti.id IS NULL OR expected_item.sku_key IS NOT NULL)
        """,
        (org_id, f"%{MANAGED_TEMPLATE_MARKER}%"),
    )
    template_count, template_item_count = cursor.fetchone()
    actual = {
        "products": product_count,
        "supplier_links": link_count,
        "templates": template_count,
        "template_items": template_item_count,
    }
    expected = {
        "products": len(products),
        "supplier_links": expected_links,
        "templates": expected_suppliers,
        "template_items": expected_links,
    }
    if actual != expected:
        raise RuntimeError(f"Post-write verification failed. Expected {expected}, got {actual}")
    return actual


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", required=True, type=Path)
    parser.add_argument("--host", default="82.25.110.112")
    parser.add_argument("--port", type=int, default=1234)
    parser.add_argument("--database", default="containermgmt_pg")
    parser.add_argument("--user", default="postgres")
    parser.add_argument("--org-id", type=int)
    parser.add_argument("--sslmode", default="require")
    parser.add_argument("--connect-timeout", type=int, default=10)
    parser.add_argument("--statement-timeout-ms", type=int, default=120000)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument(
        "--checkpoint-dir",
        default=str(Path(__file__).resolve().parents[2] / "backups"),
    )
    parser.add_argument("--checkpoint-timeout", type=int, default=900)
    parser.add_argument("--pg-dump")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    products, workbook_stats = load_source(args.excel.resolve())
    print(f"Workbook validated: {workbook_stats}")

    connection = None
    try:
        connection, password = connect(args)
        connection.autocommit = False
        with connection.cursor() as cursor:
            verify_schema(cursor)
            org_id = resolve_org_id(cursor, args.org_id, list(products))
            inspection = inspect_database(cursor, products, org_id)
            print(f"Database check passed: {inspection}")
        connection.rollback()

        if not args.apply:
            print("Dry run complete. No database changes were made.")
            return 0

        connection.close()
        connection = None
        checkpoint = create_checkpoint(args, password, org_id)
        print(f"Checkpoint created: {checkpoint} ({checkpoint.stat().st_size} bytes)")

        connection, _ = connect(args, password)
        connection.autocommit = False
        with connection.cursor() as cursor:
            verify_schema(cursor)
            resolved_org_id = resolve_org_id(cursor, org_id, list(products))
            changes = upsert_data(cursor, products, resolved_org_id)
            verified = verify_applied(cursor, products, resolved_org_id)
            print(f"Pending changes: {changes}")
            print(f"Verification passed: {verified}")
        connection.commit()
        print("Import committed successfully.")
        return 0
    except Exception as exc:
        if connection is not None and not connection.closed:
            connection.rollback()
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        if connection is not None:
            connection.close()


if __name__ == "__main__":
    raise SystemExit(main())
